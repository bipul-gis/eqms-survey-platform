import argparse
import json
import re
from pathlib import Path

from openpyxl import load_workbook


def clean(value):
    return '' if value is None else str(value).strip()


def strip_question_number(label):
    return re.sub(r'^\s*[0-9০-৯]+(?:\.[0-9০-৯]+)*\.?\s+', '', clean(label))


def split_boolean_expression(expression):
    parts = []
    depth = 0
    quote = None
    start = 0
    index = 0
    while index < len(expression):
        char = expression[index]
        if quote:
            if char == quote:
                quote = None
        elif char in ("'", '"'):
            quote = char
        elif char == '(':
            depth += 1
        elif char == ')':
            depth -= 1
        elif depth == 0:
            match = re.match(r'\s+(and|or)\s+', expression[index:], re.IGNORECASE)
            if match:
                parts.append(expression[start:index].strip())
                parts.append(match.group(1).upper())
                index += match.end()
                start = index
                continue
        index += 1
    parts.append(expression[start:].strip())
    return parts


def parse_relevant(expression):
    expression = clean(expression)
    if not expression:
        return None
    tokens = split_boolean_expression(expression)
    conditions = []
    combinators = []
    for token in tokens:
        if token in ('AND', 'OR'):
            combinators.append(token)
            continue
        selected = re.fullmatch(
            r"selected\(\$\{([^}]+)\},\s*['\"]([^'\"]+)['\"]\)", token
        )
        comparison = re.fullmatch(
            r"\$\{([^}]+)\}\s*(>=|<=|!=|=|>|<)\s*(['\"]?)(.*?)\3",
            token,
        )
        if selected:
            question_id, value = selected.groups()
            operator = 'contains'
        elif comparison:
            question_id, symbol, _, value = comparison.groups()
            operator = {
                '=': 'equals',
                '!=': 'notEquals',
                '>': 'greaterThan',
                '>=': 'greaterThanOrEqual',
                '<': 'lessThan',
                '<=': 'lessThanOrEqual',
            }[symbol]
            value = value.strip()
        else:
            return None
        conditions.append({
            'id': f'xls_logic_{len(conditions) + 1}',
            'questionId': question_id,
            'operator': operator,
            'value': value,
        })
    if len(conditions) == 1:
        combinator = 'AND'
    elif not combinators or len(set(combinators)) != 1:
        return None
    else:
        combinator = combinators[0]
    return {'enabled': True, 'combinator': combinator, 'conditions': conditions}


def constraint_validation(question_type, constraint, constraint_message):
    validation = {'integerOnly': question_type == 'integer'} if question_type in ('integer', 'decimal') else {}
    expression = clean(constraint)
    for symbol, key in (('>=', 'min'), ('<=', 'max'), ('>', 'min'), ('<', 'max')):
        match = re.search(rf"\.\s*{re.escape(symbol)}\s*(-?\d+(?:\.\d+)?)", expression)
        if match:
            number = float(match.group(1))
            validation[key] = int(number) if number.is_integer() else number
    message = clean(constraint_message)
    return validation or None, message or None


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument('workbook', type=Path)
    parser.add_argument('--output', type=Path, required=True)
    args = parser.parse_args()

    workbook = load_workbook(args.workbook, read_only=True, data_only=True)
    survey_rows = list(workbook['survey'].iter_rows(values_only=True))
    survey_headers = {name: index for index, name in enumerate(survey_rows[0])}
    choice_rows = list(workbook['choices'].iter_rows(values_only=True))
    choice_headers = {name: index for index, name in enumerate(choice_rows[0])}

    choices = {}
    for row in choice_rows[1:]:
        list_name = clean(row[choice_headers['list_name']])
        name = clean(row[choice_headers['name']])
        if not list_name or not name:
            continue
        english = clean(row[choice_headers['label::English (en)']]) or name
        bangla = clean(row[choice_headers['label::Bangla (bn)']]) or english
        choices.setdefault(list_name, []).append({
            'id': name,
            'value': name,
            'label': english,
            'labelTranslations': {'en': english, 'bn': bangla},
        })

    index = survey_headers.__getitem__
    output_questions = []
    unsupported_relevant = []
    repeat_names = []
    source_question_count = 0
    section_depth = 0

    for row_number, row in enumerate(survey_rows[1:], start=2):
        raw_type = clean(row[index('type')])
        if not raw_type:
            continue
        name = clean(row[index('name')]) or f'xls_row_{row_number}'
        english = clean(row[index('label::English (en)')])
        bangla = clean(row[index('label::Bangla (bn)')])
        translations = {'en': english, 'bn': bangla}
        if raw_type in ('end_group', 'end_repeat'):
            if raw_type == 'end_group':
                section_depth = max(0, section_depth - 1)
            continue
        if raw_type == 'begin_group':
            section_depth += 1
            output_questions.append({
                'id': f'__section__{name}',
                'type': 'section',
                'key': name,
                'question': english or bangla or name,
                'questionTranslations': translations,
                'required': False,
            })
            continue
        if raw_type == 'begin_repeat':
            repeat_names.append(name)
            output_questions.append({
                'id': f'__repeat__{name}',
                'type': 'section',
                'key': name,
                'question': english or bangla or name,
                'questionTranslations': translations,
                'description': 'Repeat record in source XLSForm. The platform currently captures one record per repeat section.',
                'descriptionTranslations': {
                    'en': 'Repeat record in source XLSForm. The platform currently captures one record per repeat section.',
                    'bn': 'মূল XLSForm-এ এই অংশে একাধিক রেকর্ড নেওয়া হয়। প্ল্যাটফর্মে আপাতত প্রতি পুনরাবৃত্ত অংশে একটি রেকর্ড নেওয়া হবে।',
                },
                'required': False,
            })
            continue
        if raw_type == 'note':
            output_questions.append({
                'id': name,
                'type': 'section',
                'key': name,
                'question': english or bangla or name,
                'questionTranslations': translations,
                'required': False,
            })
            continue

        question_english = strip_question_number(english)
        question_bangla = strip_question_number(bangla)
        question_translations = {'en': question_english, 'bn': question_bangla}
        source_question_count += 1
        appearance = clean(row[index('appearance')]).lower()
        base_type, _, list_name = raw_type.partition(' ')
        question_type = {
            'text': 'longtext' if appearance == 'multiline' else 'text',
            'integer': 'number',
            'decimal': 'number',
            'date': 'date',
            'time': 'time',
            'geopoint': 'location',
            'image': 'photo',
        }.get(base_type)
        options = None
        if base_type == 'select_one':
            question_type = 'select'
            options = choices.get(list_name, [])
        elif base_type == 'select_multiple':
            question_type = 'multiselect'
            options = choices.get(list_name, [])
        constraint_text = clean(row[index('constraint')])
        if base_type == 'text' and re.search(r'regex\(\.\s*,\s*[\'\"]\^0\[0-9\]\{10\}', constraint_text):
            question_type = 'phone'
        if not question_type:
            continue

        required_value = clean(row[index('required')]).lower()
        question = {
            'id': name,
            'key': name,
            'type': question_type,
            'question': question_english or question_bangla or name,
            'questionTranslations': question_translations,
            'required': required_value in ('true', 'yes', '1'),
        }
        hint_en = clean(row[index('hint::English (en)')])
        hint_bn = clean(row[index('hint::Bangla (bn)')])
        if hint_en or hint_bn:
            question['description'] = hint_en or hint_bn
            question['descriptionTranslations'] = {'en': hint_en, 'bn': hint_bn}
        if options is not None:
            question['options'] = options
        if base_type in ('integer', 'decimal'):
            validation, message = constraint_validation(
                base_type,
                row[index('constraint')],
                row[index('constraint_message::English (en)')],
            )
            if validation:
                question['validation'] = validation
            if message:
                question['validation']['errorMessage'] = message
        elif question_type == 'phone':
            question['validation'] = {'digits': 11}
            message = clean(row[index('constraint_message::English (en)')])
            if message:
                question['validation']['errorMessage'] = message
        relevant = clean(row[index('relevant')])
        if relevant:
            logic = parse_relevant(relevant)
            if logic:
                question['logic'] = logic
            else:
                unsupported_relevant.append((name, relevant))
        output_questions.append(question)

    payload = {
        'source': args.workbook.name,
        'sourceQuestionCount': source_question_count,
        'repeatSections': repeat_names,
        'unsupportedRelevant': unsupported_relevant,
        'questions': output_questions,
    }
    args.output.parent.mkdir(parents=True, exist_ok=True)
    args.output.write_text(json.dumps(payload, ensure_ascii=False, separators=(',', ':')) + '\n', encoding='utf-8')
    print(f'Generated {len(output_questions)} platform entries from {source_question_count} XLSForm questions.')
    print(f'Choice lists: {len(choices)}; repeat sections: {len(repeat_names)}; unsupported relevance rules: {len(unsupported_relevant)}.')
    for name, expression in unsupported_relevant:
        print(f'UNSUPPORTED RELEVANT {name}: {expression}')


if __name__ == '__main__':
    main()